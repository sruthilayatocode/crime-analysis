import docx
from docx.enum.section import WD_SECTION_START
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
# Section 1: title
doc.add_paragraph('Title')

# Section 2: 2 columns
sec2 = doc.add_section(WD_SECTION_START.CONTINUOUS)
sectPr = sec2._sectPr
cols = parse_xml(f'<w:cols {nsdecls("w")} w:num="2" w:space="360"/>')
sectPr.append(cols)
doc.add_paragraph('Col 1 paragraph')
doc.save('test_cols.docx')
print('Success!')
